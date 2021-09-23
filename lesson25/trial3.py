import os
import docx
import re
import openpyxl

def findNum(string):
    pattern = re.compile(r"\d+")
    res = re.findall(pattern, string)
    for number in res:
        lenNumber = len(number)
        if lenNumber == 11:
            return number

searchPath = "/Users/chilly/Desktop/python/yequ/再见啦/problem25"
PhoneList = []
allItems = os.listdir(searchPath)
for item in allItems:
        itemPath = os.path.join(searchPath, item)
        if os.path.isdir(itemPath):
            continue
        extension = os.path.splitext(item)[1].lower()
        if extension == ".docx":
            doc = docx.Document(itemPath)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        cellText = cell.text
                        newNum = findNum(cellText)
                        if newNum == None:
                            continue
                        PhoneList.append(newNum)

        if extension == ".xlsx":
            wb = openpyxl.load_workbook(itemPath, data_only=True)
            for ws in wb.worksheets:
                for row in ws.rows:
                    for cell in row:
                        cellValue = str(cell.value)
                        newNum = findNum(cellValue)
                        if newNum == None:
                            continue
                        PhoneList.append(newNum)

NewList = list(set(PhoneList))
print(NewList)
