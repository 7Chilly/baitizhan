import openpyxl
import docx

wb = openpyxl.load_workbook("员工花名册.xlsx")
ws = wb["汇总"]
for row in ws.rows:
    doc = docx.Document("八周年邀请函模版.docx")
    name = row[0].value
    if name == "姓名":
        continue
    doc.paragraphs[3].runs[0].add_text(name)
    doc.save(f"邀请函/八周年邀请函_致{name}.docx")
print("success")
