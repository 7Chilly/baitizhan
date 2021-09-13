import docx
import openpyxl

wb = openpyxl.load_workbook("重点单词.xlsx")
aSheet = wb["数学"]
doc = docx.Document("数学英语词汇.docx")
CorrectMeaning = []
for para in doc.paragraphs:
    for run in para.runs:
        if run.bold or run.underline:
            ChineseMeaning = para.runs[0].text
            CorrectMeaning.append(ChineseMeaning)

for idx,answer in enumerate(CorrectMeaning, 1):
    if aSheet[f"B{idx}"].value != answer:
        aSheet[f"C{idx}"].value = answer

wb.save("重点单词.xlsx")
