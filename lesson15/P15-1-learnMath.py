import docx
import openpyxl

wb = openpyxl.Workbook()
aSheet = wb["Sheet"]
aSheet.title = "数学"
doc = docx.Document("数学英语词汇.docx")
for para in doc.paragraphs:
    for run in para.runs:
        if run.bold or run.underline:
            aSheet.append([run.text])

wb.save("重点单词.xlsx")
